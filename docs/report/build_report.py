"""Build the AeroMind project report from the university template.

    python3 docs/report/build_report.py     # -> docs/AeroMind_Project_Report.docx (+ .pdf preview)

Keeps the template's cover, certificate and declaration, fills them in, and replaces the guidance
pages with the report written in content.py. Figures come from docs/report/assets (screenshots.py,
diagrams.py); numbers come from outputs/, history.json and docs/RESULTS.md. Page numbers in the
contents, figure and table lists are taken from a first rendering (LibreOffice) and are fields, so
Word refreshes them on open.
"""
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(ROOT))

import content  # noqa: E402
import docx_kit as kit  # noqa: E402

TEMPLATE = ROOT / "docs" / "Project Report Format.docx"
OUTPUT = ROOT / "docs" / "AeroMind_Project_Report.docx"
TITLE = "AeroMind: Multi-Agent Post-Flight Analysis of ArduPilot Drone Missions"
TEXT_WIDTH_IN = 6.27
FIRST_GENERATED = "ACKNOWLEDGEMENT"


class Report:
    def __init__(self, pages: dict = None):
        self.doc = docx.Document(TEMPLATE)
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.line_spacing = 1.0
        self.pages = pages or {}
        self.headings, self.figures, self.tables = [], [], []
        self._front_matter()
        kit.styles(self.doc)
        kit.page_setup(self.doc)
        self.bullet_id = None

    # -- template front matter -------------------------------------------------------------------
    def _front_matter(self) -> None:
        body = self.doc.element.body
        elements = list(body)
        cut = next(i for i, e in enumerate(elements)
                   if "".join(t.text or "" for t in e.iter(qn("w:t"))).strip() == FIRST_GENERATED)
        for e in elements[cut:]:
            if e.tag != qn("w:sectPr"):
                body.remove(e)
        paragraphs = self.doc.paragraphs
        for p in paragraphs:
            p.paragraph_format.line_spacing = 1.0          # template pages were laid out single-spaced
            if p.paragraph_format.space_after is None:
                p.paragraph_format.space_after = Pt(0)     # and without the body's paragraph spacing
            text = p.text
            if text.strip() == "“PROJECT TITLE”":
                self._replace(p, f"“{TITLE}”", keep_highlight=False)
            elif "entitled “" in text and "_____" in text:
                self._replace_blank_title(p)
            elif text.strip() == "Academic Year: 2026–27":
                pass
        # the cover: a two-line title must still fit one page
        cover_end = next(k for k, p in enumerate(paragraphs) if p.text.strip() == "CERTIFICATE")
        for p in paragraphs[:cover_end]:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if p.text.strip().startswith("“AeroMind"):
                for run in p.runs:
                    run.font.size = Pt(20)
        # the template spaces its pages with empty paragraphs sized for US Letter; use page breaks
        for name in ("CERTIFICATE", "DECLARATION"):
            p = next(p for p in self.doc.paragraphs if p.text.strip() == name)
            p.paragraph_format.page_break_before = True
        for p in list(self.doc.paragraphs):
            if not p.text.strip() and not p._p.findall(".//" + qn("w:drawing")) and p._p.getparent() is body:
                prev = p._p.getprevious()
                if prev is not None and prev.tag == qn("w:p") and not "".join(
                        t.text or "" for t in prev.iter(qn("w:t"))).strip():
                    body.remove(p._p)

    @staticmethod
    def _replace(paragraph, text: str, keep_highlight: bool = True) -> None:
        runs = paragraph.runs
        runs[0].text = text
        if not keep_highlight:
            runs[0].font.highlight_color = None
        for run in runs[1:]:
            run._r.getparent().remove(run._r)

    @staticmethod
    def _replace_blank_title(paragraph) -> None:
        for run in paragraph.runs:
            if "_____" in run.text:
                run.text = re.sub(r"_{5,}", TITLE, run.text)

    # -- headings and text -----------------------------------------------------------------------
    def h1(self, text: str, numbered: bool = True) -> None:
        p = self.doc.add_paragraph(text, style="Heading 1")
        p.paragraph_format.page_break_before = True
        self.headings.append((1, text))

    def h2(self, text: str) -> None:
        self.doc.add_paragraph(text, style="Heading 2")
        self.headings.append((2, text))

    def h3(self, text: str) -> None:
        self.doc.add_paragraph(text, style="Heading 3")

    def p(self, text: str, align=None, bold: bool = False):
        para = self.doc.add_paragraph()
        kit.inline(para, text)
        if bold:
            for run in para.runs:
                run.bold = True
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return para

    def bullets(self, items: list) -> None:
        self._list(items, "bullet", "")

    def numbered(self, items: list) -> None:
        self._list(items, "decimal", "%1.")

    def _list(self, items: list, fmt: str, text: str) -> None:
        num_id = kit._numbering(self.doc, fmt, text)
        for item in items:
            para = self.doc.add_paragraph()
            para.paragraph_format.space_after = 0
            kit.inline(para, item)
            kit.set_list(para, num_id)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(6)

    def code(self, text: str, caption: str = None) -> None:
        lines = text.strip("\n").splitlines()
        table = self.doc.add_table(rows=1, cols=1)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        kit.shade(cell, "F3F5F7")
        cell.paragraphs[0].style = self.doc.styles["Code"]
        cell.paragraphs[0].text = lines[0]
        for line in lines[1:]:
            cell.add_paragraph(line, style="Code")
        kit.fix_widths(table, [TEXT_WIDTH_IN])
        if caption:
            self.p(caption).runs[0].italic = True
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # -- numbered figures and tables ---------------------------------------------------------------
    def figure(self, path: Path, caption: str, width_in: float = TEXT_WIDTH_IN) -> int:
        number = len(self.figures) + 1
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.line_spacing = 1.0
        para.add_run().add_picture(str(path), width=Inches(width_in))
        cap = self.doc.add_paragraph(style="Caption")
        cap.add_run("Figure ").bold = True
        kit.add_field(cap, "SEQ Figure \\* ARABIC", str(number)).bold = True
        cap.add_run(f": {caption}").bold = False
        self.figures.append((number, caption))
        return number

    def table(self, headers: list, rows: list, caption: str, widths: list = None) -> int:
        number = len(self.tables) + 1
        cap = self.doc.add_paragraph(style="Caption")
        cap.paragraph_format.keep_with_next = True
        cap.add_run("Table ").bold = True
        kit.add_field(cap, "SEQ Table \\* ARABIC", str(number)).bold = True
        cap.add_run(f": {caption}")
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, header in enumerate(headers):
            cell = table.rows[0].cells[k]
            cell.paragraphs[0].style = self.doc.styles["Table Text"]
            kit.inline(cell.paragraphs[0], f"**{header}**")
            kit.shade(cell, "E3E8EE")
        kit.repeat_header(table.rows[0])
        for row in rows:
            cells = table.add_row().cells
            for k, value in enumerate(row):
                cells[k].paragraphs[0].style = self.doc.styles["Table Text"]
                kit.inline(cells[k].paragraphs[0], str(value))
        widths = widths or [1] * len(headers)
        total = sum(widths)
        kit.fix_widths(table, [TEXT_WIDTH_IN * w / total for w in widths])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        self.tables.append((number, caption))
        return number

    # -- contents, lists of figures and tables -------------------------------------------------------
    def contents(self) -> None:
        self.h1("TABLE OF CONTENTS", numbered=False)
        self._toc_marker = len(self.doc.element.body) - 2      # the heading (the last element is sectPr)

    def lists(self) -> None:
        self.h1("LIST OF FIGURES", numbered=False)
        self._lof_marker = len(self.doc.element.body) - 2
        self.h1("LIST OF TABLES", numbered=False)
        self._lot_marker = len(self.doc.element.body) - 2

    def finish(self) -> None:
        """Insert the three field lists (now that every heading, figure and table is known)."""
        body = self.doc.element.body
        marker_elements = [body[self._toc_marker], body[self._lof_marker], body[self._lot_marker]]
        page = lambda key: self.pages.get(key, "")
        toc = [(text, page(("h", text))) for level, text in self.headings
               if text not in ("TABLE OF CONTENTS",)]
        lof = [(f"Figure {n}: {c}", page(("f", n))) for n, c in self.figures]
        lot = [(f"Table {n}: {c}", page(("t", n))) for n, c in self.tables]
        levels = dict((text, level) for level, text in self.headings)
        for anchor, (instruction, entries, style) in zip(marker_elements, (
                ('TOC \\o "1-2" \\h \\z \\u', toc, None),
                ('TOC \\h \\z \\c "Figure"', lof, "Table of Figures"),
                ('TOC \\h \\z \\c "Table"', lot, "Table of Figures"))):
            start = len(body) - 1
            if style is None:
                entries = [(t, p_, f"TOC {levels.get(t, 1)}") for t, p_ in entries]
            kit.field_list(self.doc, instruction, entries, style or "TOC 1")
            created = list(body)[start:len(body) - 1]
            for e in reversed(created):
                anchor.addnext(e)
        kit.update_fields_on_open(self.doc)

    def save(self, path: Path) -> None:
        self.finish()
        self.doc.save(path)


def render_pdf(path: Path, out_dir: Path) -> Path:
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(path)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    return out_dir / (path.stem + ".pdf")


def page_texts(pdf: Path) -> list:
    count = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True,
                                                             text=True).stdout).group(1))
    texts = []
    for page in range(1, count + 1):
        out = subprocess.run(["pdftotext", "-f", str(page), "-l", str(page), "-layout", str(pdf), "-"],
                             capture_output=True, text=True).stdout
        texts.append(" ".join(out.split()))
    return texts


def find_pages(report: Report, texts: list) -> dict:
    """Page of each heading, figure and table caption, searching after the front lists."""
    start = max(k for k, t in enumerate(texts) if "LIST OF TABLES" in t) + 1
    pages = {}
    for _, text in report.headings:
        key = " ".join(text.split())
        front = key in ("ACKNOWLEDGEMENT", "ABSTRACT", "TABLE OF CONTENTS", "LIST OF FIGURES", "LIST OF TABLES")
        # a front-matter heading is the first thing on its page (its name also appears in the contents)
        found = next((k for k in range(len(texts)) if (texts[k].strip().startswith(key) if front
                                                       else key in texts[k] and k >= start)), None)
        if found is not None:
            pages[("h", text)] = str(found + 1)
    for kind, items in (("f", report.figures), ("t", report.tables)):
        label = "Figure" if kind == "f" else "Table"
        for n, _ in items:
            found = next((k for k in range(start, len(texts)) if f"{label} {n}:" in texts[k]), None)
            if found is not None:
                pages[(kind, n)] = str(found + 1)
    return pages


def build(pages: dict = None) -> Report:
    report = Report(pages)
    content.write(report)
    return report


def main() -> int:
    work = Path(tempfile.mkdtemp(prefix="aeromind-report-"))
    first = build({("h", "x"): "00"})
    draft = work / OUTPUT.name
    first.save(draft)
    texts = page_texts(render_pdf(draft, work))
    final = build(find_pages(first, texts))
    final.save(OUTPUT)
    pdf = render_pdf(OUTPUT, OUTPUT.parent)
    texts = page_texts(pdf)
    appendix = next((k for k, t in enumerate(texts) if t.startswith("APPENDIX") or " APPENDIX " in t[:40]), len(texts))
    print(f"wrote {OUTPUT.name}: {len(texts)} pages ({appendix} before the appendix); preview {pdf.name}")
    shutil.rmtree(work, ignore_errors=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
