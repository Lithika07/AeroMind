"""Small python-docx toolkit for the project report: styles, numbered captions, lists, code, fields."""
import copy
import re

from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

FONT = "Times New Roman"
MONO = "Courier New"
BLACK = RGBColor(0, 0, 0)
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


# -- page and styles -------------------------------------------------------------------------------

def page_setup(document) -> None:
    """A4, 1 inch margins on all sides, page number bottom centre (not on the cover)."""
    for section in document.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Inches(1))
        section.different_first_page_header_footer = True
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = ""
        add_field(footer, "PAGE", "1")


def _font(style, size, bold=False, italic=False, name=FONT):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = BLACK
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(theme), None)


def styles(document) -> None:
    s = document.styles
    normal = s["Normal"]
    _font(normal, 12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)
    for name, size, before, after in (("Heading 1", 16, 0, 18), ("Heading 2", 14, 14, 6), ("Heading 3", 12, 10, 4)):
        style = s[name]
        _font(style, size, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == "Heading 1" else WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    caption = s.add_style("Caption", 1) if "Caption" not in [x.name for x in s] else s["Caption"]
    _font(caption, 11, italic=False)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    code = s.add_style("Code", 1)
    _font(code, 9, name=MONO)
    code.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Inches(0.2)
    cell = s.add_style("Table Text", 1)
    _font(cell, 10.5)
    cell.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cell.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell.paragraph_format.space_after = Pt(0)
    for level, size in ((1, 12), (2, 12)):
        toc = s.add_style(f"TOC {level}", 1)
        _font(toc, size, bold=(level == 1))
        toc.paragraph_format.left_indent = Inches(0.3 * (level - 1))
        toc.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        toc.paragraph_format.space_after = Pt(3)
        toc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tof = s.add_style("Table of Figures", 1)
    _font(tof, 12)
    tof.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    tof.paragraph_format.space_after = Pt(3)
    tof.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


# -- lists -------------------------------------------------------------------------------------------

def _numbering(document, fmt: str, text: str) -> int:
    """A new list definition (bullet or decimal); returns its numId."""
    numbering = document.part.numbering_part.element
    abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    aid, nid = max(abstract_ids, default=0) + 1, max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(aid))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", fmt), ("w:lvlText", text), ("w:lvlJc", "left")):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        lvl.append(e)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    lvl.append(ppr)
    if fmt == "bullet":
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi"):
            fonts.set(qn(a), "Symbol")
        rpr.append(fonts)
        lvl.append(rpr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(nid))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(aid))
    num.append(ref)
    numbering.append(num)
    return nid


def set_list(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(nid)
    before = {qn(t) for t in ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl")}
    position = sum(1 for child in ppr if child.tag in before)
    ppr.insert(position, numpr)


# -- fields ------------------------------------------------------------------------------------------

def _fld(run, kind: str) -> None:
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), kind)
    run._r.append(e)


def _instr(run, text: str) -> None:
    e = OxmlElement("w:instrText")
    e.set(qn("xml:space"), "preserve")
    e.text = f" {text} "
    run._r.append(e)


def add_field(paragraph, instruction: str, cached: str):
    """A field with a cached result (shown until the reader's word processor updates it)."""
    _fld(paragraph.add_run(), "begin")
    _instr(paragraph.add_run(), instruction)
    _fld(paragraph.add_run(), "separate")
    result = paragraph.add_run(cached)
    _fld(paragraph.add_run(), "end")
    return result


def field_list(document, instruction: str, entries: list, style: str) -> None:
    """A TOC-type field (contents, figures, tables) whose cached result is `entries`
    [(text, page)] or [(text, page, style)]."""
    entries = entries or [("(update field)", "")]
    for k, entry in enumerate(entries):
        text, page = entry[0], entry[1]
        p = document.add_paragraph(style=entry[2] if len(entry) > 2 else style)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if k == 0:
            _fld(p.add_run(), "begin")
            _instr(p.add_run(), instruction)
            _fld(p.add_run(), "separate")
        p.add_run(f"{text}\t{page}")
        if k == len(entries) - 1:
            _fld(p.add_run(), "end")


def update_fields_on_open(document) -> None:
    settings = document.settings.element
    e = OxmlElement("w:updateFields")
    e.set(qn("w:val"), "true")
    later = {qn("w:" + t) for t in ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
                                    "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
                                    "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                                    "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
                                    "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")}
    anchor = next((child for child in settings if child.tag in later), None)
    if anchor is not None:
        anchor.addprevious(e)
    else:
        settings.append(e)


# -- content helpers ---------------------------------------------------------------------------------

def inline(paragraph, text: str) -> None:
    """Text with **bold** and `code` spans."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO
            run.font.size = Pt(10.5)
        else:
            paragraph.add_run(part)


def page_break(document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def shade(cell, hex_fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def keep_row_together(row) -> None:
    """Never split the row across pages (cantSplit precedes tblHeader in CT_TrPr)."""
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    trpr.insert(0, e)


def repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trpr.append(e)


def fix_widths(table, widths_in: list) -> None:
    """Column widths that Word and LibreOffice both honour: grid, cells, table width, fixed layout."""
    tbl = table._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    tblw.set(qn("w:type"), "dxa")
    head = {qn(t) for t in ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
                            "w:tblStyleColBandSize")}
    tblpr.insert(sum(1 for child in tblpr if child.tag in head), tblw)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tail = [child for child in tblpr if child.tag in {qn(t) for t in ("w:tblCellMar", "w:tblLook", "w:tblCaption",
                                                                       "w:tblDescription")}]
    if tail:
        tail[0].addprevious(layout)
    else:
        tblpr.append(layout)
    grid = tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for w in widths_in:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(w * 1440)))
        grid.append(col)
    for row in table.rows:
        for k, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[k])


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True
